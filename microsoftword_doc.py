from docx import Document

document = Document()
document.add_paragraph("It was a dark and stormy night.")
document.save("dark-and-stormy.docx")
document = Document("dark-and-stormy.docx")
print(document.paragraphs[0].text)
